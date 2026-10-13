"""Offline OCR and durable local storage. No network is used by offline engines."""
from __future__ import annotations
import base64
import ctypes
import io
import json
import math
import os
import re
import sqlite3
import sys
import threading
import time
import uuid
from pathlib import Path
from urllib.parse import urlparse

import numpy as np
from PIL import Image, ImageOps

BASE = Path(getattr(sys, '_MEIPASS', Path(__file__).parent))
APP_VERSION = '1.1.0'

class Cancelled(Exception):
    pass

def load_image(source):
    if isinstance(source, Image.Image):
        im = source.copy()
    else:
        im = Image.open(source)
    im = ImageOps.exif_transpose(im)
    if im.width * im.height > 40_000_000:
        raise ValueError('图片超过 4000 万像素，请先裁剪后重试。')
    if im.mode in ('RGBA', 'LA') or 'transparency' in im.info:
        rgba = im.convert('RGBA')
        bg = Image.new('RGBA', im.size, 'white')
        bg.alpha_composite(rgba)
        im = bg.convert('RGB')
    else:
        im = im.convert('RGB')
    return im

def normalize_background(im):
    a = np.asarray(im.convert('L'))
    border = np.concatenate([a[0],a[-1],a[:,0],a[:,-1]])
    if np.median(border) < 110:
        return ImageOps.invert(im)
    return im

def image_data(im):
    buf = io.BytesIO()
    im.save(buf, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode('ascii')

def session(path, threads=4):
    import onnxruntime as ort
    opts = ort.SessionOptions()
    opts.intra_op_num_threads = threads
    opts.inter_op_num_threads = 1
    opts.log_severity_level = 3
    return ort.InferenceSession(str(path), opts, providers=['CPUExecutionProvider'])

class FormulaEngine:
    def __init__(self):
        from tokenizers import Tokenizer
        root = BASE/'models'/'formula'
        self.encoder = session(root/'encoder_model.onnx')
        self.decoder = session(root/'decoder_model.onnx')
        self.tokenizer = Tokenizer.from_file(str(root/'tokenizer.json'))

    def recognize(self, im, cancel, progress=lambda s: None):
        im = normalize_background(im)
        # The upstream DeiT processor resizes directly to 384 x 384 with bicubic.
        a = np.asarray(im.resize((384,384),Image.Resampling.BICUBIC),dtype=np.float32)
        pixels = ((a / 255.0 - .5) / .5).transpose(2,0,1)[None].copy()
        hidden = self.encoder.run(None, {'pixel_values': pixels})[0]
        ids = [1]
        scores = []
        for step in range(768):
            if cancel.is_set():
                raise Cancelled()
            values = {'input_ids': np.array([ids],dtype=np.int64), 'encoder_hidden_states': hidden}
            logits = self.decoder.run(None, values)[0][0,-1].astype(np.float64)
            logits[0] = logits[1] = -1e9
            token = int(logits.argmax())
            p = np.exp(logits - np.max(logits))
            scores.append(float(p[token]/p.sum()))
            if token == 2:
                break
            ids.append(token)
            if step and step % 40 == 0:
                progress(f'正在识别公式 · {step} 个符号')
            if len(ids)>80 and len(set(ids[-24:])) <= 2:
                raise ValueError('模型出现重复输出。请只框选一个公式，并保持图片清晰。')
        else:
            raise ValueError('公式过长，已达到识别长度限制。请拆分为较小的区域。')
        text = self.tokenizer.decode(ids, skip_special_tokens=True).strip()
        if not text:
            raise ValueError('没有识别到公式，请调整截图区域。')
        return {'content':text, 'kind':'formula', 'confidence': float(np.mean(scores)) if scores else None}

class OfflineEngine:
    def __init__(self):
        self.formula = None
        self.text = None
        self.layout = None
        self.lock = threading.Lock()

    def formula_ready(self, progress):
        # The lock lets the start-up warm-up and the first recognition share one load.
        with self.lock:
            if self.formula is None:
                progress('正在加载离线公式模型…')
                self.formula = FormulaEngine()
        return self.formula

    def warm(self):
        try: self.formula_ready(lambda s: None)
        except Exception: pass

    def text_ready(self, progress):
        if self.text is None:
            progress('正在加载离线文字模型…')
            from rapidocr_onnxruntime import RapidOCR
            self.text = RapidOCR(intra_op_num_threads=4, inter_op_num_threads=1)
        return self.text

    def detect_formula(self, im, progress):
        if self.layout is None:
            progress('正在加载公式定位模型…')
            self.layout = session(BASE/'models'/'layout'/'pix2text-mfd-1.5.onnx')
        import cv2
        size = 768
        input_meta = self.layout.get_inputs()[0]
        if isinstance(input_meta.shape[-1], int):
            size = input_meta.shape[-1]
        ratio = min(size/im.width, size/im.height)
        w,h = round(im.width*ratio), round(im.height*ratio)
        left,top = (size-w)//2,(size-h)//2
        canvas = Image.new('RGB',(size,size),(114,114,114))
        canvas.paste(im.resize((w,h),Image.Resampling.BILINEAR),(left,top))
        a = np.asarray(canvas,dtype=np.float32).transpose(2,0,1)[None]/255
        pred = self.layout.run(None,{input_meta.name:a})[0]
        pred = np.squeeze(pred)
        if pred.shape[0]<pred.shape[1]:
            pred=pred.T
        scores=pred[:,4:].max(axis=1)
        cls=pred[:,4:].argmax(axis=1)
        keep=scores>.35
        pred,scores,cls=pred[keep],scores[keep],cls[keep]
        boxes=[]
        for x,y,bw,bh,*_ in pred:
            boxes.append([float(x-bw/2),float(y-bh/2),float(bw),float(bh)])
        indices=cv2.dnn.NMSBoxes(boxes,scores.tolist(),.35,.45)
        found=[]
        for i in np.asarray(indices).flatten():
            x,y,bw,bh=boxes[i]
            rect=(max(0,int((x-left)/ratio)-3),max(0,int((y-top)/ratio)-3),min(im.width,int((x+bw-left)/ratio)+3),min(im.height,int((y+bh-top)/ratio)+3))
            if rect[2]>rect[0] and rect[3]>rect[1]:
                found.append((rect,int(cls[i]),float(scores[i])))
        return found

    def recognize(self, im, mode, cancel, progress):
        im=normalize_background(im)
        if mode=='formula':
            return self.formula_ready(progress).recognize(im,cancel,progress)
        if mode=='text':
            progress('正在识别中英文文字…')
            result,_=self.text_ready(progress)(np.asarray(im)[:,:,::-1].copy())
            if cancel.is_set(): raise Cancelled()
            if not result: raise ValueError('没有识别到文字。')
            return {'content':'\n'.join(r[1] for r in result),'kind':'text','confidence':float(np.mean([r[2] for r in result]))}
        # Locate formula regions, mask them for text OCR, then merge in reading order.
        from PIL import ImageDraw
        boxes=self.detect_formula(im,progress)
        regions=[]
        mask=im.copy()
        draw=ImageDraw.Draw(mask)
        for number,(rect,kind,score) in enumerate(boxes):
            if cancel.is_set(): raise Cancelled()
            progress(f'识别图文中的公式 {number+1}/{len(boxes)}…')
            result=self.formula_ready(progress).recognize(im.crop(rect),cancel,progress)
            # Isolated scalar labels can be misread as a superscript with no base.
            result['content']=re.sub(r'^\^\s*\{\s*([a-zA-Z])\s*\}$',r'\1',result['content'])
            regions.append({'rect':rect,'text':r'\(' + result['content'] + r'\)','score':result['confidence'],'math':True})
            draw.rectangle(rect,fill='white')
        result,_=self.text_ready(progress)(np.asarray(mask)[:,:,::-1].copy())
        if cancel.is_set(): raise Cancelled()
        for box,text,score in result or []:
            a=np.asarray(box)
            rect=(float(a[:,0].min()),float(a[:,1].min()),float(a[:,0].max()),float(a[:,1].max()))
            math_like=not re.search(r'[\u3400-\u9fff]',text) and (bool(re.search(r'[=≥≤∫∑√]',text)) or (text.lower().startswith(('min','max')) and len(text)>4))
            if math_like:
                x,y,x2,y2=rect
                crop=im.crop((max(0,int(x)-3),max(0,int(y)-4),min(im.width,int(x2)+3),min(im.height,int(y2)+4)))
                converted=self.formula_ready(progress).recognize(crop,cancel,progress)
                text=r'\('+converted['content']+r'\)'
                score=converted['confidence']
            regions.append({'rect':rect,'text':text,'score':float(score),'math':math_like})
        if not regions: raise ValueError('没有识别到文字或公式，请调整截图区域。')
        regions.sort(key=lambda r:(r['rect'][1],r['rect'][0]))
        lines=[]
        for r in regions:
            x,y,x2,y2=r['rect']
            cy=(y+y2)/2
            line=next((line for line in reversed(lines[-4:]) if max(line['height'],y2-y)/max(1,min(line['height'],y2-y))<2.5 and abs(line['cy']-cy)<.45*max(16,min(line['height'],y2-y))),None)
            if line is None:
                lines.append({'cy':cy,'height':y2-y,'items':[r]})
            else: line['items'].append(r)
        content='\n\n'.join(' '.join(r['text'] for r in sorted(line['items'],key=lambda r:r['rect'][0])) for line in sorted(lines,key=lambda l:min(r['rect'][1] for r in l['items'])))
        return {'content':content,'kind':'mixed','confidence':float(np.mean([r['score'] for r in regions]))}

def cloud_recognize(im, settings, mode):
    import requests
    data_url=image_data(im)
    if settings['provider']=='mathpix':
        if not settings.get('mathpix_id') or not settings.get('api_key'):
            raise ValueError('请在设置中填写 Mathpix App ID 和 API Key。')
        response=requests.post('https://api.mathpix.com/v3/text',headers={'app_id':settings['mathpix_id'],'app_key':settings['api_key']},json={'src':data_url,'formats':['text','latex_styled'],'data_options':{'include_latex':True}},timeout=(15,120))
        if response.status_code!=200:
            raise ValueError(f'Mathpix 请求失败（HTTP {response.status_code}），请检查密钥、额度和网络。')
        data=response.json()
        content=data.get('latex_styled') if mode=='formula' else data.get('text')
        if not content: content=data.get('text')
        if not content: raise ValueError('Mathpix 没有返回可识别内容。')
        return {'content':strip_fence(content),'kind':'formula' if mode=='formula' else 'mixed','confidence':data.get('confidence')}
    endpoint=settings.get('endpoint','').strip()
    parsed=urlparse(endpoint)
    if parsed.scheme!='https' and not (parsed.scheme=='http' and parsed.hostname in ('localhost','127.0.0.1','::1')):
        raise ValueError('接口地址必须使用 HTTPS（本机 localhost 可用 HTTP）。')
    if not settings.get('model'): raise ValueError('请填写支持图片输入的模型名称。')
    prompt={'formula':'Transcribe the mathematical formula in this image into LaTeX. Return only raw LaTeX, without dollar signs, explanations or code fences. Do not solve it.','text':'Transcribe all text in the image in reading order. Return only the original text. Preserve Chinese and English. Do not explain.','mixed':'Transcribe this image into Markdown. Preserve all Chinese and English text. Use \\( \\) for inline formulas and \\[ \\] for display formulas. Return only the transcription, no code fences. Do not solve or explain.'}[mode]
    headers={'Content-Type':'application/json'}
    if settings.get('api_key'): headers['Authorization']='Bearer '+settings['api_key']
    response=requests.post(endpoint,headers=headers,json={'model':settings['model'],'messages':[{'role':'user','content':[{'type':'text','text':prompt},{'type':'image_url','image_url':{'url':data_url}}]}],'max_tokens':4096,'temperature':0},timeout=(15,120))
    if response.status_code!=200: raise ValueError(f'接口请求失败（HTTP {response.status_code}），请检查地址、模型、密钥与额度。')
    data=response.json()
    choice=data.get('choices',[{}])[0]
    if choice.get('finish_reason')=='length': raise ValueError('联网输出被截断，请缩小识别区域。')
    content=choice.get('message',{}).get('content','')
    if not isinstance(content,str) or not content.strip(): raise ValueError('接口没有返回文本，请使用支持图片输入的 Chat Completions 接口。')
    return {'content':strip_fence(content),'kind':'formula' if mode=='formula' else mode,'confidence':None}

def strip_fence(text):
    text=re.sub(r'^```[^\n]*\n|\n```\s*$','',text.strip())
    for l,r in [('$$','$$'),(r'\[',r'\]'),(r'\(',r'\)'),('$','$')]:
        if text.startswith(l) and text.endswith(r):
            text=text[len(l):-len(r)].strip()
            break
    return text

_TOKEN=re.compile(r'\\[a-zA-Z]+|\\.|\s+|.',re.S)
# Arguments of these commands are copied verbatim: spaces inside them are meaningful.
_VERBATIM={'\\text','\\textbf','\\textit','\\textrm','\\textsf','\\texttt','\\mbox','\\hbox','\\begin','\\end'}
_SIZED={'\\left','\\right','\\middle','\\big','\\Big','\\bigg','\\Bigg','\\bigl','\\bigr','\\bigm','\\Bigl','\\Bigr','\\Bigm','\\biggl','\\biggr','\\biggm','\\Biggl','\\Biggr','\\Biggm'}
_TALL={'\\frac','\\dfrac','\\tfrac','\\cfrac','\\binom','\\dbinom','\\tbinom','\\sum','\\prod','\\coprod','\\int','\\iint','\\iiint','\\oint','\\bigcup','\\bigcap','\\bigoplus','\\bigotimes','\\begin','\\underbrace','\\overbrace','\\substack','\\overset','\\underset','\\stackrel'}

def _is_word(token):
    return token.startswith('\\') and token[1:2].isalpha()

def tidy_latex(source):
    """Rewrite spaced-out model output into compact, conventional LaTeX.

    The recognition model emits one token per symbol, padded with spaces and
    stray negative thin spaces. The meaning is unchanged: redundant spaces and
    \\! are dropped, single-character scripts lose their braces, and brackets
    around tall content such as fractions become \\left ... \\right.
    """
    raw=_TOKEN.findall(source)
    tokens=[];k=0
    while k<len(raw):
        t=raw[k];k+=1
        if t.isspace() or t=='\\!':continue
        tokens.append(t)
        if t in _VERBATIM:
            j=k
            while j<len(raw) and raw[j].isspace():j+=1
            if j<len(raw) and raw[j]=='{':
                depth=0
                for m in range(j,len(raw)):
                    if raw[m]=='{':depth+=1
                    elif raw[m]=='}':
                        depth-=1
                        if depth==0:
                            tokens.append(''.join(raw[j:m+1]));k=m+1
                            break
    # Thin spaces the model puts beside operators are rendering noise, not content.
    kept=[]
    for k,t in enumerate(tokens):
        beside=[tokens[k-1] if k else '=',tokens[k+1] if k+1<len(tokens) else '=']
        if t=='\\:' or (t in ('\\,','\\;') and any(x in ('+','-','=','<','>') for x in beside)):continue
        kept.append(t)
    tokens=kept
    # x_{0} -> x_0, unless the next symbol would visually run into the script.
    out=[];k=0
    while k<len(tokens):
        body=tokens[k+2] if tokens[k] in ('_','^') and tokens[k+1:k+2]==['{'] and tokens[k+3:k+4]==['}'] else ''
        follower=tokens[k+4] if k+4<len(tokens) else ' '
        if len(body)==1 and body.isascii() and body.isalnum() and not follower[0].isalnum():
            out+=[tokens[k],body];k+=4
        else:
            out.append(tokens[k]);k+=1
    tokens=out
    # Pair brackets inside one brace group or environment cell; grow the pairs holding tall content.
    grow=set();frames=[['top',[]]]
    for k,t in enumerate(tokens):
        opens=frames[-1][1]
        prev=tokens[k-1] if k else ''
        if t=='{':frames.append(['brace',[]])
        elif t=='}':
            if len(frames)>1 and frames[-1][0]=='brace':frames.pop()
        elif t=='\\begin':frames.append(['env',[]])
        elif t=='\\end':
            while len(frames)>1 and frames.pop()[0]!='env':pass
        elif t in ('&','\\\\'):opens.clear()
        elif t in ('(','['):
            fixed=prev in _SIZED or (t=='[' and (prev in ('\\sqrt','\\\\') or (k>1 and tokens[k-2]=='\\begin')))
            opens.append((k,t,fixed))
        elif t in (')',']') and opens:
            start,kind,fixed=opens.pop()
            if kind+t in ('()','[]') and not fixed and prev not in _SIZED and any(x in _TALL for x in tokens[start+1:k]):
                grow.update((start,k))
    out=[]
    for k,t in enumerate(tokens):
        if k in grow:out.append('\\left' if t in ('(','[') else '\\right')
        out.append(t)
    parts=[];prev=''
    for t in out:
        if prev:
            after_word=_is_word(prev) and not prev.endswith('}')
            if after_word and _is_word(t):space=prev not in _SIZED
            elif after_word:space=t[0].isalnum()
            elif _is_word(t):space=t!='\\right' and (prev[-1].isalnum() or prev[-1] in "})]'")
            else:space=False
            if space or '\\\\' in (prev,t) or '&' in (prev,t):parts.append(' ')
        parts.append(t);prev=t
    return ''.join(parts)

def tidy_content(content, kind):
    if kind=='formula':return tidy_latex(content)
    if kind!='mixed':return content
    for left,right in ((r'\(',r'\)'),(r'\[',r'\]'),('$$','$$')):
        pattern=re.escape(left)+'(.+?)'+re.escape(right)
        content=re.sub(pattern,lambda m:left+tidy_latex(m.group(1))+right,content,flags=re.S)
    return content

WRAPPERS={'latex':('',''),'inline':('$','$'),'display':('$$\n','\n$$'),'bracket':('\\[\n','\n\\]'),'equation':('\\begin{equation}\n','\n\\end{equation}')}

def wrap_formula(content, kind, style):
    if kind!='formula':return content
    left,right=WRAPPERS.get(style,WRAPPERS['latex'])
    return left+content+right

def protect_secret(text, decrypt=False):
    if not text: return ''
    from ctypes import wintypes
    class Blob(ctypes.Structure):
        _fields_=[('cbData',wintypes.DWORD),('pbData',ctypes.POINTER(ctypes.c_ubyte))]
    raw=base64.b64decode(text) if decrypt else text.encode('utf-8')
    buffer=ctypes.create_string_buffer(raw)
    source=Blob(len(raw),ctypes.cast(buffer,ctypes.POINTER(ctypes.c_ubyte)))
    out=Blob()
    crypt32=ctypes.WinDLL('crypt32',use_last_error=True)
    fn=crypt32.CryptUnprotectData if decrypt else crypt32.CryptProtectData
    if not fn(ctypes.byref(source),None,None,None,None,1,ctypes.byref(out)):
        code=ctypes.get_last_error()
        raise ValueError(f'Windows 密钥加密失败（{code}），请在当前 Windows 用户下重新保存密钥。')
    try:
        value=ctypes.string_at(out.pbData,out.cbData)
        return value.decode('utf-8') if decrypt else base64.b64encode(value).decode('ascii')
    finally:
        ctypes.windll.kernel32.LocalFree(out.pbData)

DEFAULTS={'provider':'offline','mode':'formula','auto_copy':True,'copy_format':'latex','tidy':True,'theme':'auto','close_to_tray':True,'hotkey':'Ctrl+Alt+S','endpoint':'','model':'','mathpix_id':'','key_offline':'','key_mathpix':'','key_compatible':''}

class Store:
    def __init__(self, root):
        self.root=Path(root)
        self.root.mkdir(parents=True,exist_ok=True)
        (self.root/'images').mkdir(exist_ok=True)
        self.db=sqlite3.connect(self.root/'history.sqlite')
        self.db.row_factory=sqlite3.Row
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.execute('CREATE TABLE IF NOT EXISTS snips(id TEXT PRIMARY KEY, created TEXT, title TEXT, content TEXT, kind TEXT, image TEXT, confidence REAL, engine TEXT, elapsed REAL, favorite INTEGER DEFAULT 0)')
        self.db.execute('CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT)')
        self.db.commit()

    def settings(self, secret=False):
        cfg=DEFAULTS.copy()
        cfg.update({r['key']:json.loads(r['value']) for r in self.db.execute('SELECT * FROM settings')})
        encrypted=cfg.get('key_'+cfg['provider'],'')
        if secret:
            cfg['api_key']=protect_secret(encrypted,True) if encrypted else ''
        else:
            cfg={k:v for k,v in cfg.items() if not k.startswith('key_')}
            cfg['has_key']=bool(encrypted)
        return cfg

    def save_settings(self, data):
        provider=data.get('provider','offline')
        if provider not in ('offline','mathpix','compatible'): raise ValueError('未知引擎')
        for key,val in data.items():
            if key in DEFAULTS and not key.startswith('key_'):
                self.db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)',(key,json.dumps(val)))
        if data.get('api_key'):
            self.db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)',('key_'+provider,json.dumps(protect_secret(data['api_key']))))
        if data.get('clear_key'):
            self.db.execute('DELETE FROM settings WHERE key=?',('key_'+provider,))
        self.db.commit()

    def put(self, key, value):
        self.db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)',(key,json.dumps(value)))
        self.db.commit()

    def add(self, result, im=None, title='未命名公式', engine='手动编辑', elapsed=0):
        from datetime import datetime
        ident=uuid.uuid4().hex
        image=''
        if im is not None:
            image=ident+'.png'
            im.save(self.root/'images'/image)
        self.db.execute('INSERT INTO snips VALUES (?,?,?,?,?,?,?,?,?,0)',(ident,datetime.now().astimezone().isoformat(timespec='seconds'),title,result['content'],result['kind'],image,result.get('confidence'),engine,elapsed))
        self.db.commit()
        return ident

    def rows(self):
        return [dict(r) for r in self.db.execute('SELECT * FROM snips ORDER BY created DESC, rowid DESC')]

    def get(self, ident):
        row=self.db.execute('SELECT * FROM snips WHERE id=?',(ident,)).fetchone()
        if not row: raise ValueError('记录不存在')
        return dict(row)

    def image(self, row, thumb=False):
        if not row.get('image'): return ''
        try:
            im=load_image(self.root/'images'/Path(row['image']).name)
            if thumb: im.thumbnail((460,160))
            return image_data(im)
        except (OSError,ValueError): return ''

    def update(self, ident, content, title, kind):
        if kind not in ('formula','text','mixed'): raise ValueError('格式无效')
        self.db.execute('UPDATE snips SET content=?,title=?,kind=?,confidence=NULL WHERE id=?',(content,title[:200],kind,ident))
        self.db.commit()

    def delete(self, ident):
        row=self.get(ident)
        self.db.execute('DELETE FROM snips WHERE id=?',(ident,))
        self.db.commit()
        if row['image']:
            (self.root/'images'/Path(row['image']).name).unlink(missing_ok=True)

    def favorite(self, ident):
        self.db.execute('UPDATE snips SET favorite=1-favorite WHERE id=?',(ident,))
        self.db.commit()

def formats(content, kind):
    if kind!='formula': return [('Markdown / 文本',content)]
    return [('LaTeX',content),('行内公式',r'\('+content+r'\)'),('独立公式',r'\['+'\n'+content+'\n'+r'\]'),('$$ 块公式','$$\n'+content+'\n$$'),('带编号公式',r'\begin{equation}'+'\n'+content+'\n'+r'\end{equation}')]

def mathml(content):
    import latex2mathml.converter
    return latex2mathml.converter.convert(content)

def export_docx(row, path):
    from docx import Document
    from docx.oxml import parse_xml
    from docx.shared import Pt
    import mathml2omml
    document=Document()
    document.styles['Normal'].font.name='Calibri'
    document.styles['Normal'].font.size=Pt(12)
    document.add_heading(row['title'],0)
    def equation(source):
        omml=mathml2omml.convert(mathml(source))
        p=document.add_paragraph()
        wrapper=parse_xml('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'+omml+'</w:p>')
        for node in list(wrapper):p._p.append(node)
    if row['kind']=='formula':
        equation(row['content'])
    elif row['kind']=='mixed':
        parts=re.split(r'(\\\[.*?\\\]|\\\(.*?\\\)|\$\$.*?\$\$)',row['content'],flags=re.S)
        for part in parts:
            if not part: continue
            if part.startswith((r'\[',r'\(','$$')): equation(part[2:-2])
            else:
                for line in part.strip().splitlines(): document.add_paragraph(line)
    else:
        for line in row['content'].splitlines(): document.add_paragraph(line)
    document.save(str(path))
